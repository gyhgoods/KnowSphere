import csv
import json
from dataclasses import dataclass
from io import BytesIO, StringIO
from pathlib import Path

from docx import Document as WordDocument
from openpyxl import load_workbook
from pypdf import PdfReader

from app.services.text_processing import normalize_text


class UnsupportedDocumentTypeError(ValueError):
    pass


TEXT_SUFFIXES = {".txt", ".md", ".markdown", ".log", ".xml", ".yaml", ".yml"}
IMAGE_RELATION_MARKER = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"


@dataclass(frozen=True, slots=True)
class ParsedImage:
    index: int
    file_name: str
    mime_type: str
    content: bytes
    description: str


@dataclass(frozen=True, slots=True)
class ParsedDocument:
    text: str
    images: list[ParsedImage]


def clean_text(parts: list[str]) -> str:
    return normalize_text("\n".join(part for part in parts if part))


def decode_text(content: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="replace")


def parse_pdf(content: bytes) -> ParsedDocument:
    reader = PdfReader(BytesIO(content))
    return ParsedDocument(
        text=clean_text([page.extract_text() or "" for page in reader.pages]),
        images=[],
    )


def parse_word(content: bytes) -> ParsedDocument:
    document = WordDocument(BytesIO(content))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append("\t".join(cell.text.strip() for cell in row.cells))

    images: list[ParsedImage] = []
    for relation in document.part.rels.values():
        if relation.reltype != IMAGE_RELATION_MARKER:
            continue
        target = relation.target_part
        image_index = len(images) + 1
        suffix = Path(str(target.partname)).suffix or ".bin"
        file_name = f"word-image-{image_index}{suffix}"
        title = getattr(target, "partname", file_name)
        images.append(
            ParsedImage(
                index=image_index,
                file_name=file_name,
                mime_type=target.content_type,
                content=target.blob,
                description=(
                    f"Word embedded image {image_index}. Source part: {title}. "
                    "This image was extracted from the document for visual reference."
                ),
            )
        )
    return ParsedDocument(text=clean_text(parts), images=images)


def parse_excel(content: bytes) -> ParsedDocument:
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    parts: list[str] = []
    try:
        for worksheet in workbook.worksheets:
            parts.append(f"[{worksheet.title}]")
            for row in worksheet.iter_rows(values_only=True):
                values = ["" if value is None else str(value) for value in row]
                if any(values):
                    parts.append("\t".join(values))
    finally:
        workbook.close()
    return ParsedDocument(text=clean_text(parts), images=[])


def parse_csv(content: bytes) -> ParsedDocument:
    text = decode_text(content)
    return ParsedDocument(
        text=clean_text(["\t".join(row) for row in csv.reader(StringIO(text))]),
        images=[],
    )


def parse_json(content: bytes) -> ParsedDocument:
    value = json.loads(decode_text(content))
    return ParsedDocument(text=json.dumps(value, ensure_ascii=False, indent=2), images=[])


def parse_document_assets(file_name: str, mime_type: str, content: bytes) -> ParsedDocument:
    suffix = Path(file_name).suffix.lower()
    if suffix == ".pdf" or mime_type == "application/pdf":
        return parse_pdf(content)
    if suffix == ".docx":
        return parse_word(content)
    if suffix == ".xlsx":
        return parse_excel(content)
    if suffix == ".csv" or mime_type == "text/csv":
        return parse_csv(content)
    if suffix == ".json" or mime_type == "application/json":
        return parse_json(content)
    if suffix in TEXT_SUFFIXES or mime_type.startswith("text/"):
        return ParsedDocument(text=clean_text(decode_text(content).splitlines()), images=[])
    raise UnsupportedDocumentTypeError(f"Unsupported document type: {suffix or mime_type}")


def parse_document(file_name: str, mime_type: str, content: bytes) -> str:
    return parse_document_assets(file_name, mime_type, content).text
