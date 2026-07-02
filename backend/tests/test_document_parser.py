import base64
from io import BytesIO

import pytest
from docx import Document as WordDocument
from openpyxl import Workbook
from pypdf import PdfWriter

from app.services.document_parser import (
    UnsupportedDocumentTypeError,
    parse_document,
    parse_document_assets,
)
from app.tasks.document_index_task import extract_image_chunks

TINY_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMB"
    "/gL+X8sAAAAASUVORK5CYII="
)


def test_parse_plain_text_csv_and_json() -> None:
    assert parse_document("guide.txt", "text/plain", "知识 文档".encode()) == "知识 文档"
    assert parse_document("data.csv", "text/csv", b"name,value\nalpha,1") == (
        "name\tvalue\nalpha\t1"
    )
    assert '"name": "KnowSphere"' in parse_document(
        "data.json", "application/json", b'{"name":"KnowSphere"}'
    )


def test_parse_word_document() -> None:
    stream = BytesIO()
    document = WordDocument()
    document.add_heading("KnowSphere", level=1)
    document.add_paragraph("Enterprise knowledge")
    document.save(stream)

    result = parse_document(
        "guide.docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        stream.getvalue(),
    )
    assert "KnowSphere" in result
    assert "Enterprise knowledge" in result


def test_parse_word_document_images() -> None:
    stream = BytesIO()
    document = WordDocument()
    document.add_paragraph("Image section")
    document.add_picture(BytesIO(TINY_PNG))
    document.save(stream)

    result = parse_document_assets(
        "guide.docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        stream.getvalue(),
    )

    assert "Image section" in result.text
    assert len(result.images) == 1
    assert result.images[0].mime_type == "image/png"


def test_extract_image_chunks_from_parse_markers() -> None:
    text, images = extract_image_chunks(
        "Before\n"
        "[[KNOWSPHERE_IMAGE index=1 object=documents/1/extracted-images/2/a.png "
        "name=word-image-1.png mime=image/png]]\n"
        "Image 1: Word embedded image 1.\n"
        "[[END_KNOWSPHERE_IMAGE]]\n"
        "After"
    )

    assert "Before" in text
    assert "After" in text
    assert len(images) == 1
    assert images[0].object_name.endswith("/a.png")
    assert "Word embedded image" in images[0].content


def test_parse_excel_workbook() -> None:
    stream = BytesIO()
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Knowledge"
    worksheet.append(["title", "status"])
    worksheet.append(["Guide", "published"])
    workbook.save(stream)
    workbook.close()

    result = parse_document(
        "guide.xlsx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        stream.getvalue(),
    )
    assert "[Knowledge]" in result
    assert "Guide\tpublished" in result


def test_parse_pdf_document() -> None:
    stream = BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.write(stream)

    assert parse_document("guide.pdf", "application/pdf", stream.getvalue()) == ""


def test_reject_unsupported_document() -> None:
    with pytest.raises(UnsupportedDocumentTypeError):
        parse_document("archive.zip", "application/zip", b"not a document")
